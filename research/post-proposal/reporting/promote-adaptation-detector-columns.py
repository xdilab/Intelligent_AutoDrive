#!/usr/bin/env python3
"""Promote six verified adaptation detector results without rewriting workbook assets."""
from pathlib import Path
from datetime import datetime
import argparse,hashlib,io,json,re,zipfile
from lxml import etree as E
import openpyxl

def sha(b):return hashlib.sha256(b).hexdigest()
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def q(n):return '{'+NS+'}'+n

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--target',type=Path,required=True);ap.add_argument('--results',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);a=ap.parse_args()
 assert a.target.resolve()!=a.output.resolve()
 raw=a.target.read_bytes();w=openpyxl.load_workbook(io.BytesIO(raw));s=w.worksheets[0]
 assert s.title=='Clean Metrics Fall 2026'
 assert [s.cell(1,i).value for i in range(13,21)]==['agentness','agent','action','loc','duplex','triplet','Tail mAP (47, z<0)','notes']
 expected={};records=[];candidate=set();frames=set()
 for stage in ['stage5','stage6']:
  for condition,variant in [('frozen','Frozen encoder + focal heads'),('classification','ROAD focal only'),('contrastive','ROAD focal + contrastive')]:
   source=a.results/(('epoch1-final-' if condition=='frozen' else 'final-')+stage+'-'+condition+'.json');d=json.loads(source.read_text())
   assert (d['stage'],d['condition'],d['protocol']['seed'],d['n_frames'])==(stage,condition,0,36717)
   assert d['selected']==('epoch-1' if condition=='frozen' else 'epoch-3')
   assert {k:v['classes'] for k,v in d['tail'].items()}=={'tail47':47,'deep28':28,'common39':39}
   candidate.add(d['candidate_sha256']);frames.add(d['frame_sha256'])
   matches=[i for i in range(2,s.max_row+1) if str(s.cell(i,2).value).startswith('Stage '+stage[-1]+':') and str(s.cell(i,3).value).split('\n')[0]==variant and str(s.cell(i,4).value).startswith('ROAD full adaptation;')]
   assert len(matches)==1;row=matches[0];assert s.cell(row,9).value=='crop AP (%)'
   old=s.cell(row,20).value;epoch=re.search(r'Epoch (?:dev triplet )?crop AP=\[[^\]]+\]',old);assert epoch
   values={'I':'f-mAP (%)','J':'YOLO candidates; locked 36,717 frames','K':'val','L':.5}
   for col,head in zip('MNOPQR',['agentness','agent','action','loc','duplex','triplet']):values[col]=d['summary'][head]
   values['S']=d['tail']['tail47']['mAP']
   values['T']=(f'Completed detector evaluation; main columns are detector f-mAP@0.5 (%). n=1; seed0 (not a three-seed mean); selected {d["selected"]} using internal development crop AP. '
     f'Preserved development triplet crop AP={s.cell(row,18).value:.12g} (%); {epoch.group(0)}. '
     f'deep28={d["tail"]["deep28"]["mAP"]:.9f}; common39={d["tail"]["common39"]["mAP"]:.9f}. '
     'Tail47: z<0; deep28: z<-0.5; common39: z>=0. '
     f'Sources: artifacts/stage56-full/collected/{stage}-{condition}.json (development); artifacts/stage56-full/collected/{source.name} (detector).')
   for col,val in values.items():expected[f'{col}{row}']=val
   records.append({'model_number':s.cell(row,1).value,'row':row,'source':str(source),'source_sha256':sha(source.read_bytes()),'previous_cells':{f'{c}{row}':s[f'{c}{row}'].value for c in values}})
 assert len(candidate)==len(frames)==1
 with zipfile.ZipFile(io.BytesIO(raw)) as z:infos=z.infolist();parts={i.filename:z.read(i.filename) for i in infos};comment=z.comment
 book=E.fromstring(parts['xl/workbook.xml']);rid=book.find(q('sheets'))[0].get('{'+REL+'}id');rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);target=next(r.get('Target') for r in rels if r.get('Id')==rid);sheet=target.lstrip('/') if target.startswith('/') else 'xl/'+target
 tree=E.fromstring(parts[sheet]);cells={c.get('r'):c for c in tree.iter(q('c'))};rows={int(r.get('r')):r for r in tree.iter(q('row'))}
 for ref,val in expected.items():
  cell=cells.get(ref)
  if cell is None:
   cell=E.Element(q('c'),r=ref);rr=rows[s[ref].row];rr.append(cell)
   rr[:]=sorted(rr,key=lambda c:openpyxl.utils.column_index_from_string(re.match('[A-Z]+',c.get('r')).group()))
  for child in list(cell):cell.remove(child)
  if isinstance(val,str):cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,q('is')),q('t')).text=val
  else:cell.set('t','n');E.SubElement(cell,q('v')).text=repr(val)
 parts[sheet]=E.tostring(tree)
 a.output.parent.mkdir(parents=True,exist_ok=True);backup=a.output.parent/'backups'/('ROAD-Waymo-before-detector-columns-'+datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'.xlsx');backup.parent.mkdir(exist_ok=True);backup.write_bytes(raw)
 with zipfile.ZipFile(a.output,'w') as z:
  z.comment=comment
  for info in infos:z.writestr(info,parts[info.filename])
 with zipfile.ZipFile(io.BytesIO(raw)) as before,zipfile.ZipFile(a.output) as after:
  assert after.testzip() is None and before.namelist()==after.namelist()
  changed=[n for n in before.namelist() if before.read(n)!=after.read(n)];assert changed==[sheet]
 nw=openpyxl.load_workbook(a.output);ns=nw.worksheets[0];assert nw.sheetnames==w.sheetnames and ns.max_row==s.max_row and ns.max_column==s.max_column
 for rr in s:
  for c in rr:
   assert ns[c.coordinate].value==expected.get(c.coordinate,c.value),c.coordinate
   assert ns[c.coordinate]._style==c._style,c.coordinate
 for i in range(1,s.max_row+1):assert dict(s.row_dimensions[i])==dict(ns.row_dimensions[i])
 audit={'time':datetime.now().astimezone().isoformat(),'target':str(a.target),'target_sha256':sha(raw),'output':str(a.output),'output_sha256':sha(a.output.read_bytes()),'backup':str(backup),'changes':expected,'sources':records,'changed_parts':changed,'preservation':'Only six rows I:T updated. All other cell values/styles and row heights checked; all other ZIP parts including drawings are byte-identical.'}
 a.output.with_suffix('.audit.json').write_text(json.dumps(audit,indent=2));print(json.dumps({'models':[r['model_number'] for r in records],'output':str(a.output),'backup':str(backup),'verification':'passed'}))
if __name__=='__main__':main()
