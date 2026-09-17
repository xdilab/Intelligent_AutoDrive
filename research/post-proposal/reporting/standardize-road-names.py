from pathlib import Path
from lxml import etree as E
import openpyxl,zipfile,datetime,json,hashlib,os
p=Path('/home/brandon/Downloads/ROAD-Waymo Results.xlsx');W=Path('/data/repos/wiki');w=openpyxl.load_workbook(p);s=w.worksheets[0];N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships';C='http://schemas.openxmlformats.org/package/2006/content-types';q=lambda t:'{'+N+'}'+t
stage={0:'Stage 0 — 3D-RetinaNet (I3D)',1:'Stage 1 — YOLO + I3D Score Transfer',2:'Stage 2 — YOLO + I3D RoI Classifier',3:'Stage 3 — I3D Composition MLP',4:'Stage 4 — Crop Phrase Head',5:'Stage 5 — Visual Composition Expert',6:'Stage 6 — Phrase-Fusion MLP'}
fixed={3:('3D-RetinaNet (I3D-08)','Published reference'),4:('3D-RetinaNet (SlowFast-08)','Published reference'),5:('YOLOv8 Agent Detector','Published reference; agent-only'),6:('YOLOv8x/m Agent-Tube Ensemble','ECCV24 Track1 winner; tube metric'),8:(stage[0],'Local all-anchor replication'),9:(stage[0],'Gödel loss; tainted constraint arrays'),10:('Qwen-ViT + FCOS','Exp1b'),11:('Qwen-ViT + DETR','Exp2'),12:('EfficientNet + Deformable DETR','Exp2b'),13:('EfficientNet + CLIP-Fused Deformable DETR','Exp2c'),14:('Swin-L + Deformable DETR','Exp2d'),15:('ResNet-50 + Deformable DETR','Exp2e; paper resolution'),16:('ResNet-50 + Deformable DETR (184-output head)','Exp2f; flat multilabel head'),17:('MS-DETR (one-to-many supervision)','Exp2g; stopped epoch4'),19:(stage[0],'Top40-box scoring control'),20:('Qwen-VL Box Classifier','Zero-shot; JSON outputs'),21:('Qwen-VL Box Classifier','Detector-steered prompt; JSON outputs'),22:('Qwen-VL Box Classifier','Joint ROAD/BDD-X/CoVLA LoRA; JSON outputs'),23:('Detector–VLM Score Fusion','Zero-shot VLM evidence'),24:('Detector–VLM Score Fusion','Zero-language ablation'),25:('Detector–VLM Score Fusion','BDD-X LoRA VLM evidence'),26:('Detector–VLM Score Fusion','Joint LoRA VLM evidence'),27:('Qwen-ViT RoI Classifier','R1; ROAD supervision'),28:('Qwen-ViT RoI Classifier','R2; corrected Gödel lambda0.1'),29:('Qwen-ViT RoI Classifier','R2; corrected Gödel lambda1'),30:('Qwen-ViT RoI Classifier','R2; corrected Gödel lambda10'),31:('Qwen-ViT RoI Classifier','R3; ROAD + language co-training'),32:('Qwen-ViT RoI Classifier','R4; co-training + corrected Gödel lambda10'),33:('InternVideo2 Feature-Map Flat Head','Exp10 C1; top40 protocol'),34:('InternVideo2 Feature-Map Phrase Head','Exp10 C2; top40 protocol'),36:(stage[0],'Own top300 detections; matched frames'),37:('YOLOv8x Agent Detector','Our detector; epoch1'),38:('YOLO26x Agent Detector','One-to-one output; epoch2'),39:(stage[1],'IoU-matched score transfer'),40:(stage[2],'Focal only; confidence gated'),41:(stage[2],'Gödel lambda10; confidence gated'),42:(stage[2],'YOLO junk negatives; confidence gated'),43:('I3D RoI Classifier + Derived Composition','Minimum composition; no composition training'),44:('I3D Composition MLP (scores only)','Primitive49 inputs; no feature input'),45:(stage[3],'Primitive49 + I3D256 inputs'),46:(stage[2],'Batch16K; focal only; confidence gated'),48:('InternVideo2 Feature-Map Flat Head','Static repeated frames; Exp12 C1'),49:('InternVideo2 Feature-Map Phrase Head','Static repeated frames; Exp12 C2'),50:('I3D RoI Phrase Head','I3D256 projected to phrase space'),51:('I3D Phrase-Primitive Composition MLP','Phrase primitives + I3D features'),53:('Crop Flat Head','InternVideo2; probe-trained; full validation'),54:(stage[4],'Probe-trained; full validation'),55:('Crop Flat Head','Full crop training'),56:(stage[4],'Full crop training'),57:(stage[5],'Original single-run study'),58:('Crop Phrase-Primitive Composition MLP','Phrase primitives + crop features; NOT Stage6 fusion'),59:('Crop Phrase-Primitive Composition MLP','Existing duplicate of preceding row; retained'),60:(stage[4],'Deterministic phrase vocabulary v2'),61:(stage[6],'Original single-run study'),62:('Crop Flat Head','Coordinate-probe paired control; 2/12 training shards'),63:('Crop Flat Head + Box Coordinates','Coordinate probe; adds cx,cy,w,h')}
condition={'frozen':'Frozen visual encoder; focal head training','classification':'Final visual block adapted; focal loss','contrastive':'Final visual block adapted; focal +0.001 contrastive'}
changes=[]
for ri in range(3,s.max_row+1):
 old=s.cell(ri,2).value
 if s.cell(ri,1).value is None:continue
 if ri in fixed:name,detail=fixed[ri]
 elif 65<=ri<=92:
  study,rest=old.split(': ',1);enc='BDD-X-adapted encoder' if rest.endswith('(adapted)') else 'Original encoder';detail=f'{study}; {enc}'
  if 'shuffled-class-gate' in rest:name='Stage 5 + Stage 4 Classwise Gate';detail+='; shuffled phrase control; BCE fitting'
  elif 'class-gate' in rest:name='Stage 5 + Stage 4 Classwise Gate';detail+='; BCE fitting'
  elif 'shuffled-global' in rest:name='Stage 5 + Stage 4 Global Blend';detail+='; shuffled phrase control; BCE fitting'
  elif 'global-blend' in rest:name='Stage 5 + Stage 4 Global Blend';detail+='; BCE fitting'
  elif 'ap-shuffled-blend' in rest:name='Stage 5 + Stage 4 Global Blend';detail+='; shuffled phrase control; development AP selection'
  elif 'ap-blend' in rest:name='Stage 5 + Stage 4 Global Blend';detail+='; development AP selection'
  elif 'Flat-evidence' in rest:name='Crop Flat-Evidence Fusion MLP';detail+='; matched-width flat-score control'
  elif 'shuffled-phrase' in rest:name=stage[6];detail+='; shuffled phrase evidence'
  elif 'Shuffled phrase head' in rest:name=stage[4];detail+='; shuffled class-to-phrase mapping'
  elif 'Flat crop head' in rest:name='Crop Flat Head'
  elif 'Stage 4' in rest:name=stage[4]
  elif 'Stage 5' in rest:name=stage[5]
  elif 'Stage 6' in rest:name=stage[6]
  else:
   import re
   num=int(re.search(r'stage([0-3])',rest)[1]);name=stage[num]
 elif 94<=ri<=113:
  study,rest=old.split(': ',1);arch,cond,ck=rest.split(' / ');name=stage[int(arch[-1])] if arch in ['stage5','stage6'] else 'Crop Flat Head'
  detail=f'{study}; {condition[cond]}; {ck}'
  if study=='ROAD small pilot':detail+='; simplified pilot, NOT Stage5/6'
  if ck=='baseline':detail+='; initial checkpoint before this arm updates'
 elif 115<=ri<=118:name='InternVideo2 Video–Text Dual Encoder';detail='BDD-X contrastive subset pilot; '+old.rsplit(' ',1)[1]
 elif ri in [120,121]:name='FRCB (MViT)';detail='Published Table5; five-run mean' if ri==120 else 'Published Table3; best-result row'
 else:raise AssertionError((ri,old))
 desc=detail
 if ri<=63 and s.cell(ri,3).value:desc+=' | '+s.cell(ri,3).value
 notes=str(s.cell(ri,19).value or '')+' | Previous label: '+old
 changes.append((ri,name,desc,notes,old))
raw=p.read_bytes();stamp=datetime.datetime.now().strftime('%Y%m%d-%H%M%S');out=W/'artifacts/results-master';backup=out/f'ROAD-Waymo-Results-before-names-{stamp}.xlsx';backup.write_bytes(raw)
with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
root=E.fromstring(parts['xl/worksheets/sheet1.xml']);rows={int(x.get('r')):x for x in root.find(q('sheetData'))}
def setcell(ri,col,text):
 c=next(c for c in rows[ri] if c.get('r')==f'{col}{ri}')
 for el in list(c):c.remove(el)
 c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=text
for ri,name,desc,notes,old in changes:
 setcell(ri,'B',name);setcell(ri,'C',desc);setcell(ri,'S',notes)
setcell(1,'B','architecture');setcell(1,'C','study / condition / variant')
table=E.fromstring(parts['xl/tables/table1.xml']);tc=table.find(q('tableColumns'));tc[1].set('name','architecture');tc[2].set('name','study / condition / variant')
key=[['Architecture','Definition / distinguishing inputs','Conventions'],*[ [stage[i],d,'Stage number is architecture lineage, not training condition.'] for i,d in enumerate(['3D-RetinaNet with I3D; retains own detector boxes.','YOLO boxes receive scores copied from matched I3D detections.','YOLO boxes + RoIAligned I3D features + learned flat classifier.','I3D composition MLP receives49 primitive scores +256-D I3D features.','InternVideo2 crop features projected to fixed phrase embeddings; phrase head without composition MLP.','Flat49 primitive scores +1024-D crop feature -> composition MLP; no explicit phrase-score inference branch.','Flat49 primitive +135 phrase-composition scores +1024-D crop feature -> composition MLP.'])],['Crop Flat Head','InternVideo2 crop features -> Linear184; no composition MLP.','Includes simplified ROAD pilot; never label that pilot Stage5.'],['Crop Phrase-Primitive Composition MLP','Phrase primitives + crop features -> composition MLP.','Different from Stage6, which explicitly adds phrase composition scores to flat primitives.'],['Stage 5 + Stage 4 Global Blend','Scalar score mixture of visual composition and phrase experts.','BCE-fit versus development-AP-selected scalar are conditions, not new architectures.'],['Stage 5 + Stage 4 Classwise Gate','Class-dependent score mixture of visual composition and phrase experts.','Shuffled language is a control condition.'],['Feature-Map heads versus Crop heads','Feature-map extraction from full images differs from per-box crop encoding.','Do not merge these architecture names.'],['Adaptation conditions','Frozen; focal-only final-block adaptation; focal +0.001 contrastive final-block adaptation.','ROAD text encoder frozen; BDD-X pilot adapted last2 blocks of both towers.'],['Checkpoints and studies','Keep baseline/epoch and Original controlled/BDD-X paired/BCE gates/AP blend/ROAD pilot/ROAD full in columnC.','Do not compare different crop/detector protocols as identical.'],['Provenance','Previous complete labels preserved in notes; row numbers and experiment IDs retained.','No metric, tail, data-split, epoch or numerical result changed by this naming edit.']]
key+=[[name,'See first-tab variant and previous label for full experiment definition.','Legacy non-stage architecture.'] for name in sorted({c[1] for c in changes}) if name not in {r[0] for r in key}]
ws=E.Element(q('worksheet'),nsmap={None:N});cols=E.SubElement(ws,q('cols'))
for i,width in enumerate([49,105,105],1):E.SubElement(cols,q('col'),min=str(i),max=str(i),width=str(width),customWidth='1')
sd=E.SubElement(ws,q('sheetData'))
for i,vals in enumerate(key,1):
 row=E.SubElement(sd,q('row'),r=str(i))
 for j,v in enumerate(vals,1):
  c=E.SubElement(row,q('c'),r=f'{openpyxl.utils.get_column_letter(j)}{i}',t='inlineStr',s=str(s.cell(1,j).style_id if i==1 else 0));E.SubElement(E.SubElement(c,q('is')),q('t')).text=v
wb=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);ct=E.fromstring(parts['[Content_Types].xml']);sl=wb.find(q('sheets'));assert 'Architecture key' not in [x.get('name') for x in sl];sid=max(int(x.get('sheetId')) for x in sl)+1
E.SubElement(sl,q('sheet'),name='Architecture key',sheetId=str(sid),attrib={'{'+R+'}id':'rIdArchitectureKey'});E.SubElement(rels,'{'+P+'}Relationship',Id='rIdArchitectureKey',Type=R+'/worksheet',Target='worksheets/architecture-key.xml');E.SubElement(ct,'{'+C+'}Override',PartName='/xl/worksheets/architecture-key.xml',ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
updates={'xl/worksheets/sheet1.xml':E.tostring(root),'xl/tables/table1.xml':E.tostring(table),'xl/worksheets/architecture-key.xml':E.tostring(ws),'xl/workbook.xml':E.tostring(wb),'xl/_rels/workbook.xml.rels':E.tostring(rels),'[Content_Types].xml':E.tostring(ct)}
tmp=p.with_name('ROAD-Waymo Results.updating.xlsx')
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in {**parts,**updates}.items():z.writestr(n,b)
a=openpyxl.load_workbook(tmp)
for row in s:
 for c in row:
  if c.column not in [2,3,19]:assert a.worksheets[0].cell(c.row,c.column).value==c.value,c.coordinate
for ri,name,desc,notes,old in changes:assert a.worksheets[0].cell(ri,2).value==name and old in a.worksheets[0].cell(ri,19).value
for old in w.worksheets[1:]:assert list(old.values)==list(a[old.title].values) and len(old._images)==len(a[old.title]._images)
with zipfile.ZipFile(tmp) as z:
 for n,b in parts.items():
  if n not in updates:assert z.read(n)==b
assert p.read_bytes()==raw;os.replace(tmp,p)
manifest={'file':str(p),'backup':str(backup),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'changed_result_rows':len(changes),'mapping':[{'row':ri,'old':old,'architecture':name,'variant':desc} for ri,name,desc,notes,old in changes],'validation':'All non-name/non-description/non-notes cells exact; all existing other sheets and images preserved.'};(out/f'names-{stamp}.json').write_text(json.dumps(manifest,indent=2)+'\n');print('Updated',len(changes),'result rows; architecture key added; numerical cells unchanged.')
