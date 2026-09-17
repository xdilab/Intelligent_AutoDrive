from pathlib import Path
from copy import deepcopy
from lxml import etree as E
import openpyxl,zipfile,datetime,json,hashlib,os
p=Path('/home/brandon/Downloads/ROAD-Waymo Results.xlsx');w=openpyxl.load_workbook(p);s=w.worksheets[0];title='Supporting - Crop Flat Heads';assert title not in w.sheetnames
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships';C='http://schemas.openxmlformats.org/package/2006/content-types';q=lambda x:'{'+N+'}'+x
move=[i for i in range(2,s.max_row+1) if str(s.cell(i,2).value or '').startswith('Crop Flat Head')];keep=[i for i in range(1,s.max_row+1) if i not in move];assert len(move)==18
raw=p.read_bytes();stamp=datetime.datetime.now().strftime('%Y%m%d-%H%M%S');out=Path('/data/repos/wiki/artifacts/results-master');backup=out/f'ROAD-Waymo-Results-before-archive-{stamp}.xlsx';backup.write_bytes(raw)
with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
r=E.fromstring(parts['xl/worksheets/sheet1.xml']);data=r.find(q('sheetData'));old={int(row.get('r')):deepcopy(row) for row in data}
def renumber(row,i):
 row.set('r',str(i))
 for c in row:c.set('r',''.join(filter(str.isalpha,c.get('r')))+str(i))
 return row
archive=E.Element(q('worksheet'),nsmap={None:N});E.SubElement(archive,q('dimension'),ref=f'A1:T{len(move)+1}');views=E.SubElement(archive,q('sheetViews'));view=E.SubElement(views,q('sheetView'),workbookViewId='0');E.SubElement(view,q('pane'),ySplit='1',topLeftCell='A2',activePane='bottomLeft',state='frozen');archive.append(deepcopy(r.find(q('sheetFormatPr'))));archive.append(deepcopy(r.find(q('cols'))));ad=E.SubElement(archive,q('sheetData'))
for i,source in enumerate([1]+move,1):ad.append(renumber(deepcopy(old[source]),i))
E.SubElement(archive,q('autoFilter'),ref=f'A1:T{len(move)+1}')
data[:]=[]
for i,source in enumerate(keep,1):data.append(renumber(deepcopy(old[source]),i))
r.find(q('dimension')).set('ref',f'A1:T{len(keep)}');table=E.fromstring(parts['xl/tables/table1.xml']);table.set('ref',f'A1:T{len(keep)}');table.find(q('autoFilter')).set('ref',f'A1:T{len(keep)}')
wb=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);ct=E.fromstring(parts['[Content_Types].xml']);sl=wb.find(q('sheets'));sid=max(int(x.get('sheetId')) for x in sl)+1;rid='rIdSupportingCropFlat';part='xl/worksheets/supporting-crop-flat.xml';E.SubElement(sl,q('sheet'),name=title,sheetId=str(sid),attrib={'{'+R+'}id':rid});E.SubElement(rels,'{'+P+'}Relationship',Id=rid,Type=R+'/worksheet',Target='worksheets/supporting-crop-flat.xml');E.SubElement(ct,'{'+C+'}Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
updates={'xl/worksheets/sheet1.xml':E.tostring(r),'xl/tables/table1.xml':E.tostring(table),part:E.tostring(archive),'xl/workbook.xml':E.tostring(wb),'xl/_rels/workbook.xml.rels':E.tostring(rels),'[Content_Types].xml':E.tostring(ct)};tmp=p.with_name('ROAD-Waymo Results.updating.xlsx')
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in {**parts,**updates}.items():z.writestr(n,b)
a=openpyxl.load_workbook(tmp)
for name,order in [(s.title,keep),(title,[1]+move)]:
 for i,source in enumerate(order,1):
  assert [a[name].cell(i,j).value for j in range(1,21)]==[s.cell(source,j).value for j in range(1,21)],(name,i)
  assert a[name].row_dimensions[i].height==15
for orig in w.worksheets[1:]:assert list(orig.values)==list(a[orig.title].values)
with zipfile.ZipFile(tmp) as z:
 assert z.testzip() is None
 for n,b in parts.items():
  if n not in updates:assert z.read(n)==b
assert p.read_bytes()==raw;os.replace(tmp,p)
report={'file':str(p),'backup':str(backup),'moved_rows':move,'moved_experiment_ids':[s.cell(i,1).value for i in move],'archive_sheet':title,'headline_total_rows':len(keep),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'validation':'Every retained and archived cell matches original; all rows15pt; all other existing ZIP parts unchanged; no result deleted.'};(out/f'archive-crop-flat-{stamp}.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
