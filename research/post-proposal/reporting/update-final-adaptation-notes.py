#!/usr/bin/env python3
"""Refresh final Stage5/6 detector notes without rewriting crop AP or drawings."""
import argparse, hashlib, io, json, re, zipfile
from datetime import datetime
from pathlib import Path
from lxml import etree as E
import openpyxl

N = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
q = lambda tag: '{'+N+'}'+tag
sha = lambda data: hashlib.sha256(data).hexdigest()


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--target', type=Path, required=True)
    ap.add_argument('--results', type=Path, required=True)
    ap.add_argument('--output', type=Path, required=True)
    a = ap.parse_args()
    assert a.target.resolve() != a.output.resolve()
    raw = a.target.read_bytes(); wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.worksheets[0]
    assert ws.title == 'Clean Metrics Fall 2026'
    assert ws['T1'].value == 'notes' and ws['R1'].value == 'triplet'
    stamp = datetime.now().astimezone().isoformat(timespec='seconds')
    changes = []; candidate_hashes = set(); frame_hashes = set()
    for stage in ['stage5', 'stage6']:
        for condition, variant in [('classification', 'ROAD focal only'), ('contrastive', 'ROAD focal + contrastive')]:
            source = a.results/f'final-{stage}-{condition}.json'
            d = json.loads(source.read_text())
            assert d['stage'] == stage and d['condition'] == condition
            assert d['selected'] == 'epoch-3' and d['protocol']['seed'] == 0
            assert d['n_frames'] == 36717
            assert {k:v['classes'] for k,v in d['tail'].items()} == {'tail47':47,'deep28':28,'common39':39}
            candidate_hashes.add(d['candidate_sha256']); frame_hashes.add(d['frame_sha256'])
            rows = [i for i in range(2, ws.max_row+1)
                    if str(ws.cell(i,2).value).startswith(f'Stage {stage[-1]}:')
                    and str(ws.cell(i,3).value).split('\n')[0] == variant
                    and str(ws.cell(i,4).value).startswith('ROAD full adaptation;')]
            assert len(rows) == 1, (stage,condition,rows)
            row = rows[0]; old = ws.cell(row,20).value or ''
            assert ws.cell(row,9).value == 'crop AP (%)'
            epoch = re.search(r'Epoch (?:dev triplet )?crop AP=\[[^\]]+\]',old)
            assert epoch, old
            scores = ', '.join(f'{k}={d["summary"][k]:.6f}' for k in ['agentness','agent','action','loc','duplex','triplet'])
            tails = ', '.join(f'{k}={d["tail"][k]["mAP"]:.6f}' for k in ['tail47','deep28','common39'])
            new = (f'Main columns retain development crop AP (NOT detector AP). {epoch.group(0)}. '
                   f'FINAL DETECTOR f-mAP@0.5 (%): {scores}; {tails}. '
                   'Selected epoch-3; seed0 only (not a three-seed mean); 36,717 locked YOLOv8x validation frames. '
                   'Tail47: z<0; deep28: z<-0.5; common39: z>=0. '
                   f'Completed; updated {stamp}. '
                   f'Sources: artifacts/stage56-full/collected/{stage}-{condition}.json (development); '
                   f'artifacts/stage56-full/collected/{source.name} (final detector).')
            changes.append({'row':row,'model_number':ws.cell(row,1).value,'cell':f'T{row}',
                            'source':str(source),'source_sha256':sha(source.read_bytes()),'before':old,'after':new})
    assert len(candidate_hashes)==len(frame_hashes)==1
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        infos=z.infolist(); parts={i.filename:z.read(i.filename) for i in infos}; comment=z.comment
    book=E.fromstring(parts['xl/workbook.xml']); rid=book.find(q('sheets'))[0].get('{'+R+'}id')
    rels=E.fromstring(parts['xl/_rels/workbook.xml.rels'])
    target=next(r.get('Target') for r in rels if r.get('Id')==rid)
    path=target.lstrip('/') if target.startswith('/') else 'xl/'+target
    root=E.fromstring(parts[path]); cells={c.get('r'):c for c in root.iter(q('c'))}
    for change in changes:
        cell=cells[change['cell']]
        for child in list(cell): cell.remove(child)
        cell.set('t','inlineStr')
        E.SubElement(E.SubElement(cell,q('is')),q('t')).text=change['after']
    parts[path]=E.tostring(root)
    a.output.parent.mkdir(parents=True,exist_ok=True)
    backup=a.output.parent/'backups'/f'ROAD-Waymo-Results-before-final-adaptation-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx'
    backup.parent.mkdir(exist_ok=True); backup.write_bytes(raw)
    with zipfile.ZipFile(a.output,'w') as z:
        z.comment=comment
        for info in infos:z.writestr(info,parts[info.filename])
    with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(a.output) as after:
        assert after.testzip() is None and before.namelist()==after.namelist()
        changed_parts=[n for n in before.namelist() if before.read(n)!=after.read(n)]
        assert changed_parts==[path],changed_parts
    result=openpyxl.load_workbook(a.output); rs=result.worksheets[0]
    expected={c['cell']:c['after'] for c in changes}
    assert result.sheetnames==wb.sheetnames and rs.max_row==ws.max_row and rs.max_column==ws.max_column
    for row in ws:
        for cell in row:
            out=rs[cell.coordinate]
            assert out.value==expected.get(cell.coordinate,cell.value),cell.coordinate
            assert out._style==cell._style,cell.coordinate
    for i in range(1,ws.max_row+1):
        assert dict(ws.row_dimensions[i])==dict(rs.row_dimensions[i]),i
    audit={'time':stamp,'target':str(a.target),'target_sha256':sha(raw),'output':str(a.output),
           'output_sha256':sha(a.output.read_bytes()),'backup':str(backup),'changes':changes,
           'changed_parts':changed_parts,'preservation':'Only four notes cells changed; all other values/styles, row heights, drawing/media parts, sheets and table ranges preserved byte-for-byte or cell-verified.'}
    a.output.with_suffix('.audit.json').write_text(json.dumps(audit,indent=2))
    print(json.dumps({'output':str(a.output),'models':[c['model_number'] for c in changes],
                      'cells':list(expected),'changed_parts':changed_parts,'backup':str(backup)}))


if __name__=='__main__':main()
