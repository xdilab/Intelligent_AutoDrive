from pathlib import Path
import zipfile,os
from lxml import etree as E
import openpyxl
p=Path('/home/brandon/Downloads/ROAD-Waymo Results.xlsx');w=openpyxl.load_workbook(p);s=w.worksheets[0];N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda t:'{'+N+'}'+t
with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
r=E.fromstring(parts['xl/worksheets/sheet1.xml']);rows={int(x.get('r')):x for x in r.find(q('sheetData'))}
def setval(ri,col,text):
 c=next(c for c in rows[ri] if c.get('r')==f'{col}{ri}')
 for x in list(c):c.remove(x)
 c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=text
for ri in range(64,s.max_row+1):
 name=str(s.cell(ri,2).value or '')
 if name.startswith('BCE gates:') and any(t in name for t in ['class-gate','global-blend','shuffled-global']):setval(ri,'G','Experts10; gate25')
 if name.startswith('AP-selected blends:'):setval(ri,'G','Experts10; dev grid')
 if name.startswith('Original controlled:'):setval(ri,'I','Locked36,717 frames; source protocol')
parts['xl/worksheets/sheet1.xml']=E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True);tmp=p.with_name('ROAD-Waymo Results.updating.xlsx')
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
check=openpyxl.load_workbook(tmp);assert check.worksheets[0].max_column==19;os.replace(tmp,p)
