#!/usr/bin/env python3
"""Merge verified contextual rows into a downloaded results workbook without rewriting drawings.
Usage: python sync-context-workbook.py --target DOWNLOAD.xlsx --source VERIFIED.xlsx --output UPDATED.xlsx
Writes a target backup and audit JSON beside output. Source must be a verified results workbook.
Does not upload to OneDrive. Re-running against updated output does not duplicate rows.
"""
import argparse,copy,hashlib,io,json,re,zipfile
from pathlib import Path
from datetime import datetime
from lxml import etree as E
import openpyxl
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships'
q=lambda t:'{'+N+'}'+t
def main():
 a=argparse.ArgumentParser(description=__doc__);a.add_argument('--target',type=Path,required=True);a.add_argument('--source',type=Path,required=True);a.add_argument('--output',type=Path,required=True);args=a.parse_args()
 assert args.output.resolve()!=args.target.resolve(),'Write to a separate output, inspect, then replace target.'
 raw=args.target.read_bytes();sw=openpyxl.load_workbook(args.source,read_only=True);tw=openpyxl.load_workbook(io.BytesIO(raw),read_only=True)
 sv=list(sw.worksheets[0].values);tv=list(tw.worksheets[0].values);assert sv[0][:20]==tv[0][:20],'Column layout mismatch'
 start=next(i for i,r in enumerate(sv) if isinstance(r[1],str) and r[1].startswith('NEW — CONTEXTUAL RESIDUAL'))
 with zipfile.ZipFile(io.BytesIO(raw)) as z:parts={n:z.read(n) for n in z.namelist()};infos=z.infolist()
 wb=E.fromstring(parts['xl/workbook.xml']);rid=wb.find(q('sheets'))[0].get('{'+R+'}id');rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);target=next(x.get('Target') for x in rels if x.get('Id')==rid);sheetpath=target.lstrip('/') if target.startswith('/') else 'xl/'+target
 root=E.fromstring(parts[sheetpath]);data=root.find(q('sheetData'));rows={int(r.get('r')):r for r in data};template=rows[max(i for i,r in enumerate(tv,1) if isinstance(r[0],(int,float)) and not tw.worksheets[0].cell(i,2).font.bold)];styles={re.sub(r'\d','',c.get('r')):c.get('s') for c in template}
 def put(row,col,v):
  ref=col+row.get('r');c=next((x for x in row if x.get('r')==ref),None)
  if c is None:c=E.SubElement(row,q('c'),r=ref)
  if styles.get(col):c.set('s',styles[col])
  for ch in list(c):c.remove(ch)
  c.attrib.pop('t',None)
  if v is None:return
  if isinstance(v,(int,float)):E.SubElement(c,q('v')).text=str(v)
  else:c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=str(v)
 def key(r):return (r[1],r[2])
 existing={key(r):i for i,r in enumerate(tv,1) if r[1] is not None};changes=[];numeric={i for i,r in enumerate(tv,1) if isinstance(r[0],(int,float))}
 for values in sv[start:]:
  if not values[1]:continue
  k=key(values);i=existing.get(k)
  if i is None:i=max(rows)+1;row=E.SubElement(data,q('row'),r=str(i));rows[i]=row;existing[k]=i;action='add'
  else:row=rows[i];action='refresh'
  row.set('ht','15');row.set('customHeight','1')
  for col,v in zip('ABCDEFGHIJKLMNOPQRST',values[:20]):put(row,col,v)
  if isinstance(values[0],(int,float)):numeric.add(i)
  changes.append({'row':i,'action':action,'architecture':values[1]})
 for n,i in enumerate(sorted(numeric),1):put(rows[i],'A',n)
 end=max(rows);root.find(q('dimension')).set('ref',f'A1:T{end}');parts[sheetpath]=E.tostring(root)
 relpath=str(Path(sheetpath).parent/'_rels'/(Path(sheetpath).name+'.rels'));sr=E.fromstring(parts[relpath]);changed={sheetpath}
 import posixpath
 for rel in sr:
  if rel.get('Type','').endswith('/table'):
   t=rel.get('Target');path=t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(sheetpath),t));table=E.fromstring(parts[path]);table.set('ref',f'A1:T{end}');af=table.find(q('autoFilter'))
   if af is not None:af.set('ref',f'A1:T{end}')
   parts[path]=E.tostring(table);changed.add(path)
 args.output.parent.mkdir(parents=True,exist_ok=True);stamp=datetime.now().strftime('%Y%m%d-%H%M%S-%f');backup=args.output.parent/f'backup-target-{stamp}.xlsx';backup.write_bytes(raw)
 with zipfile.ZipFile(args.output,'w') as z:
  for info in infos:z.writestr(info,parts[info.filename])
 with zipfile.ZipFile(io.BytesIO(raw)) as before,zipfile.ZipFile(args.output) as after:
  assert after.testzip() is None
  assert all(before.read(n)==after.read(n) for n in before.namelist() if n not in changed)
 result=list(openpyxl.load_workbook(args.output,read_only=True).worksheets[0].values)
 for i,old in enumerate(tv):
  if i+1 not in {x['row'] for x in changes}:assert old[1:20]==result[i][1:20],f'Existing row changed: {i+1}'
 assert [result[i-1][0] for i in sorted(numeric)]==list(range(1,len(numeric)+1))
 for source_row in sv[start:]:
  if source_row[1]:assert result[existing[key(source_row)]-1][1:20]==source_row[1:20]
 report={'target':str(args.target),'target_sha256':hashlib.sha256(raw).hexdigest(),'output':str(args.output),'output_sha256':hashlib.sha256(args.output.read_bytes()).hexdigest(),'backup':str(backup),'changes':changes,'rows':end,'preserved':'All existing non-context B:T cells, other sheets, drawings, media and workbook metadata retained; contextual cells use target styles.'}
 args.output.with_suffix('.audit.json').write_text(json.dumps(report,indent=2));print(json.dumps({'rows':end,'added':sum(x['action']=='add' for x in changes),'refreshed':sum(x['action']=='refresh' for x in changes)}))
if __name__=='__main__':main()
