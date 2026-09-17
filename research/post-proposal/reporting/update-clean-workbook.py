from pathlib import Path
from copy import deepcopy
import zipfile,json,datetime,hashlib,os
from lxml import etree as E
import openpyxl
from PIL import Image
W=Path('/data/repos/wiki');target=Path('/home/brandon/Downloads/ROAD-Waymo Results.xlsx');master=Path('/home/brandon/Downloads/road_waymo_results_master.xlsx');m=openpyxl.load_workbook(master,data_only=True);before=openpyxl.load_workbook(target);s=before.worksheets[0]
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships';C='http://schemas.openxmlformats.org/package/2006/content-types';X='http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing';A='http://schemas.openxmlformats.org/drawingml/2006/main'
q=lambda t:'{'+N+'}'+t
stamp=datetime.datetime.now().strftime('%Y%m%d-%H%M%S');out=W/'artifacts/results-master';backup=out/f'ROAD-Waymo-Results-before-{stamp}.xlsx';original=target.read_bytes();backup.write_bytes(original)
with zipfile.ZipFile(target) as z:parts={n:z.read(n) for n in z.namelist()}
root=E.fromstring(parts['xl/worksheets/sheet1.xml']);data=root.find(q('sheetData'))
# Insert tail at R; move existing notes to S without rewriting any values/styles.
for row in data:
 for c in row:
  if c.get('r','').startswith('R'):c.set('r','S'+c.get('r')[1:])
cols=root.find(q('cols'))
for col in cols:
 if int(col.get('min'))>=18:col.set('min',str(int(col.get('min'))+1));col.set('max',str(int(col.get('max'))+1))
col=E.Element(q('col'),min='18',max='18',width='15',customWidth='1');cols.insert(17,col)
def cell(row,col,v,style):
 c=E.SubElement(row,q('c'),r=f'{openpyxl.utils.get_column_letter(col)}{row.get("r")}',s=str(style))
 if v is None:return c
 if isinstance(v,(int,float)):E.SubElement(c,q('v')).text=str(v)
 else:c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=str(v)
 return c
cell(data[0],18,'Tail mAP (47, z<0)',s['Q1'].style_id);data[0][:]=sorted(data[0],key=lambda c:openpyxl.utils.cell.column_index_from_string(''.join(filter(str.isalpha,c.get('r')))))
# Only two historical tail values are explicitly recorded with the same z<0 cutoff.
for ri,v in [(57,5.72),(61,5.99)]:
 row=data[ri-1];cell(row,18,v,s.cell(ri,17).style_id);row[:]=sorted(row,key=lambda c:openpyxl.utils.cell.column_index_from_string(''.join(filter(str.isalpha,c.get('r')))))
new=[];counter=66
style_row=55
styles=[s.cell(style_row,min(i,18)).style_id for i in range(1,20)];styles[17]=s.cell(style_row,17).style_id;styles[18]=s.cell(style_row,18).style_id
names={'stage5':'Stage 5 Visual Composition Expert','stage6':'Stage 6 Phrase-Fusion MLP','head-flat':'Flat crop head','head-phrase':'Stage 4 Phrase head','fusion-shuffled':'Stage 6 shuffled-phrase control','head-shuffled':'Shuffled phrase head','fusion-flat-evidence':'Flat-evidence fusion control'}
def section(text):new.append([None,text]+[None]*17)
def add(model,desc,train,enc,epochs,metric,protocol,split,iou,metrics,tail,notes,dataset='ROAD-Waymo'):
 global counter
 counter+=1;new.append([counter,model,desc,dataset,train,enc,epochs,metric,protocol,split,iou,*metrics,tail,notes])
last=None
for r in list(m['Detector summary'].values)[1:]:
 study,model,encoder,n,metric=r[:5]
 if study!=last:section('NEW — '+study.upper()+' — mean across seeds where available');last=study
 means=[r[5+2*i] for i in range(6)];tail=r[17];tail_sd=r[18];deep=r[19];common=r[21]
 desc=names.get(model,model)+'; '+encoder+' encoder'
 notes=f'n={n}; triplet SD={r[16]}; tail SD={tail_sd}; deep28={deep:.4f}; common39={common:.4f}. Sources: {r[-1]}'
 add(f'{study}: {names.get(model,model)} ({encoder})',desc,r[-2], 'BDD-X last2 blocks both towers; frozen on ROAD' if encoder=='adapted' else 'frozen encoder / trainable heads',10,'f-mAP (%)','YOLO candidates; 36,717 shared frames','val',.5,means,tail,notes)
section('NEW — ROAD DEVELOPMENT CROP AP — NOT DETECTOR mAP; TAIL NOT YET COMPUTED')
for r in list(m['Development crop AP'].values)[1:]:
 study,stage,cond,checkpoint,metric,ap,selected,scope,source=r
 add(f'{study}: {stage} / {cond} / {checkpoint}','All184 channels trained; reported metric covers86 triplets',scope,'last visual block adapted' if cond!='frozen' else 'visual encoder frozen',checkpoint,'crop AP (%)','Saved development crops','internal dev',None,[None]*5+[ap],None,f'NOT detector AP. Selected so far: {selected}. Source: {source}')
section('NEW — BDD-X PILOT RETRIEVAL — METRICS IN NOTES, NOT ROAD AP COLUMNS')
for r in list(m['BDD-X retrieval'].values)[1:]:
 dataset,ck,loss,v2t,t2v,train,dev,selected,source=r
 add('BDD-X adaptation '+ck,'Symmetric video-text contrastive pilot',f'{train} train /{dev} dev','last2 blocks both towers + projections',ck,'loss / R@1','128-video retrieval pool','internal dev',None,[None]*6,None,f'Dev loss={loss:.6f}; video-to-text R@1={v2t:.4f}%; text-to-video R@1={t2v:.4f}%; selected={selected}; {source}',dataset)
section('NEW — PUBLISHED FRCB — EXTERNAL PROTOCOL, NOT A MATCHED RERUN')
for r in list(m['Published FRCB'].values)[1:]:
 ref,ds,metric,n,agent,action,loc,duplex,triplet,sd,notes,source=r
 add(ref,'MViT visual/context + dynamic class balancing','paper600 training videos','K400 pretrained MViT',30,'f-mAP (%)','Published; different candidate protocol','val',.5,[None,agent,action,loc,duplex,triplet],None,f'{n} runs / selection; triplet SD={sd}. {notes}. {source}')
for vals in new:
 ri=len(data)+1;issection=vals[0] is None;row=E.SubElement(data,q('row'),r=str(ri),ht='44' if issection else '82',customHeight='1')
 for ci,v in enumerate(vals,1):cell(row,ci,v,s.cell(35,min(ci,18)).style_id if issection else styles[ci-1])
end=len(data);root.find(q('dimension')).set('ref',f'A1:S{end}')
table=E.fromstring(parts['xl/tables/table1.xml']);table.set('ref',f'A1:S{end}');table.find(q('autoFilter')).set('ref',f'A1:S{end}');tc=table.find(q('tableColumns'));t=deepcopy(tc[16]);t.set('id',str(max(int(c.get('id')) for c in tc)+1));t.set('name','Tail mAP (47, z<0)')
for a in list(t.attrib):
 if a.endswith('uid'):del t.attrib[a]
tc.insert(17,t);tc.set('count','19')
updates={'xl/worksheets/sheet1.xml':E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True),'xl/tables/table1.xml':E.tostring(table,xml_declaration=True,encoding='UTF-8',standalone=True)}
# Append static copies of approved diagrams, preserving all original drawing parts.
figures=[('Stage 4 aligned','artifacts/defense-30min/alignment/stage4.png','Inference architecture; GT-positive head training, YOLO candidates at evaluation.'),('Stage 5 aligned','artifacts/defense-30min/alignment/stage5.png','Visual composition expert; no inference phrase branch.'),('Stage 6 aligned','artifacts/defense-30min/alignment/stage6.png','Phrase-Fusion MLP inference architecture.'),('Stage 4 detailed','artifacts/animated-stage4-phrase/stage4-static.png','Static export of approved sequential diagram; animation remains in HTML source.'),('Stage 5 detailed','artifacts/animated-stage5-composition/stage5-static.png','Static export of approved sequential diagram; animation remains in HTML source.'),('BDD-X adaptation diagram','artifacts/animated-stage6-pilot/stage6-static.png','BDD-X pilot pretraining/adaptation/transfer diagram; NOT the current ROAD last-block study.'),('Gate architecture','artifacts/class-gate-diagrams/class-gate-architecture.png','Earlier gate-study architecture; not the current encoder-adaptation study.'),('Gate training','artifacts/class-gate-diagrams/class-gate-training.png','Earlier gate-study training protocol; not the current encoder-adaptation study.')]
wb=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);cts=E.fromstring(parts['[Content_Types].xml']);sl=wb.find(q('sheets'));sid=max(int(x.get('sheetId')) for x in sl);figmanifest=[]
for j,(name,path,note) in enumerate(figures,1):
 image=W/path;assert image.exists();b=image.read_bytes();im=Image.open(image);width=1280;height=round(width*im.height/im.width);stem=f'newdiagram{j}';part=f'xl/worksheets/{stem}.xml';drawing=f'xl/drawings/{stem}.xml';media=f'xl/media/{stem}.png'
 ws=E.Element(q('worksheet'),nsmap={None:N,'r':R});sv=E.SubElement(ws,q('sheetViews'));E.SubElement(sv,q('sheetView'),workbookViewId='0',showGridLines='0',zoomScale='70');sd=E.SubElement(ws,q('sheetData'))
 for ri,text in [(1,name),(2,note),(3,'Source: '+path)]:rr=E.SubElement(sd,q('row'),r=str(ri));cell(rr,1,text,0)
 E.SubElement(ws,q('drawing'),attrib={'{'+R+'}id':'rId1'});updates[part]=E.tostring(ws)
 sr=E.Element('{'+P+'}Relationships',nsmap={None:P});E.SubElement(sr,'{'+P+'}Relationship',Id='rId1',Type=R+'/drawing',Target='../drawings/'+stem+'.xml');updates[f'xl/worksheets/_rels/{stem}.xml.rels']=E.tostring(sr)
 dr=E.Element('{'+X+'}wsDr',nsmap={'xdr':X,'a':A});anchor=E.SubElement(dr,'{'+X+'}oneCellAnchor');fr=E.SubElement(anchor,'{'+X+'}from')
 for tag,v in [('col',0),('colOff',0),('row',4),('rowOff',0)]:E.SubElement(fr,'{'+X+'}'+tag).text=str(v)
 E.SubElement(anchor,'{'+X+'}ext',cx=str(width*9525),cy=str(height*9525));pic=E.SubElement(anchor,'{'+X+'}pic');nv=E.SubElement(pic,'{'+X+'}nvPicPr');E.SubElement(nv,'{'+X+'}cNvPr',id='1',name=name,descr=note+' Source: '+path);E.SubElement(nv,'{'+X+'}cNvPicPr');bf=E.SubElement(pic,'{'+X+'}blipFill');E.SubElement(bf,'{'+A+'}blip',attrib={'{'+R+'}embed':'rId1'});E.SubElement(E.SubElement(bf,'{'+A+'}stretch'),'{'+A+'}fillRect');sp=E.SubElement(pic,'{'+X+'}spPr');E.SubElement(sp,'{'+A+'}prstGeom',prst='rect');E.SubElement(anchor,'{'+X+'}clientData');updates[drawing]=E.tostring(dr)
 ir=E.Element('{'+P+'}Relationships',nsmap={None:P});E.SubElement(ir,'{'+P+'}Relationship',Id='rId1',Type=R+'/image',Target='../media/'+stem+'.png');updates[f'xl/drawings/_rels/{stem}.xml.rels']=E.tostring(ir);updates[media]=b
 rid='rIdNewDiagram'+str(j);E.SubElement(sl,q('sheet'),name=name,sheetId=str(sid+j),attrib={'{'+R+'}id':rid});E.SubElement(rels,'{'+P+'}Relationship',Id=rid,Type=R+'/worksheet',Target='worksheets/'+stem+'.xml')
 for p,typ in [(part,'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'),(drawing,'application/vnd.openxmlformats-officedocument.drawing+xml')]:E.SubElement(cts,'{'+C+'}Override',PartName='/'+p,ContentType=typ)
 figmanifest.append({'sheet':name,'source':path,'sha256':hashlib.sha256(b).hexdigest()})
updates.update({'xl/workbook.xml':E.tostring(wb),'xl/_rels/workbook.xml.rels':E.tostring(rels),'[Content_Types].xml':E.tostring(cts)})
tmp=target.with_name('ROAD-Waymo Results.updating.xlsx')
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in {**parts,**updates}.items():z.writestr(n,b)
after=openpyxl.load_workbook(tmp);a=after.worksheets[0]
for row in s:
 for c in row:assert a.cell(c.row,c.column if c.column<18 else c.column+1).value==c.value,(c.coordinate,'changed')
for vals,ri in zip(new,range(64,end+1)):assert [a.cell(ri,i).value for i in range(1,20)]==vals,ri
for old in before.worksheets[1:]:assert list(old.values)==list(after[old.title].values);assert len(old._images)==len(after[old.title]._images)
for name,_,_ in figures:assert len(after[name]._images)==1
assert a.tables['Table5'].ref==f'A1:S{end}' and len(a.tables['Table5'].tableColumns)==19
with zipfile.ZipFile(tmp) as z:
 for n,b in parts.items():
  if n not in updates:assert z.read(n)==b,n
assert target.read_bytes()==original;os.replace(tmp,target)
report={'file':str(target),'backup':str(backup),'timestamp':stamp,'new_result_rows':sum(x[0] is not None for x in new),'new_section_rows':sum(x[0] is None for x in new),'first_tab_total_rows':end,'tail_column':'R: Tail mAP (47, z<0); notes moved to S','aggregate_policy':'25 detector summary rows consolidate67 individual runs; seed SD and raw paths retained in notes','figures':figmanifest,'validation':'Every original cell preserved with notes shift; all other original worksheets/drawings and ZIP parts preserved; new rows and8 images read back; table range extended','sha256':hashlib.sha256(target.read_bytes()).hexdigest()};(out/f'clean-update-{stamp}.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
