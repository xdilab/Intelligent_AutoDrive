from pathlib import Path
import json,re,statistics,zipfile,hashlib,shutil,datetime,os
from collections import defaultdict
from lxml import etree as E
import openpyxl
W=Path('/data/repos/wiki');target=Path('/home/brandon/Downloads/road_waymo_results_master.xlsx')
now=datetime.datetime.now(datetime.timezone.utc).isoformat();stamp=datetime.datetime.now().strftime('%Y%m%d-%H%M%S')
out=W/'artifacts/results-master';out.mkdir(exist_ok=True)
sources={}
def read(p):
 p=Path(p);sources[str(p.relative_to(W))]=hashlib.sha256(p.read_bytes()).hexdigest();return json.loads(p.read_text())
heads=['agentness','agent','action','loc','duplex','triplet'];tails=['tail47','deep28','common39']
records=[]
for study,folder,split in [('Original controlled','post-proposal-experiments/cluster-results/metrics','full expert training'),('BDD-X paired','adapted-road-comparison/collected/metrics','paired recache, full training'),('BCE gates','class-gate-study/collected/metrics','420 expert /90 gate /90 development'),('AP-selected blends','ap-selected-blend/collected/metrics','420 expert /90 gate /90 development')]:
 for p in sorted((W/'artifacts'/folder).rglob('*.json')):
  d=read(p)
  if 'summary' not in d or 'triplet' not in d['summary']:continue
  match=re.match(r'seed(\d+)-(.+)',p.stem);seed=int(match[1]) if match else None;model=match[2] if match else p.stem;encoder=p.parent.name if study=='BDD-X paired' else 'original'
  records.append(dict(study=study,model=model,encoder=encoder,seed=seed,split=split,data=d,source=str(p.relative_to(W))))
for p in sorted((W/'artifacts/stage56-full/collected').glob('*final-*.json')):
 d=read(p);records.append(dict(study='ROAD full epoch1' if p.name.startswith('epoch1') else 'ROAD full selected',model=d['stage']+'-'+d['condition'],encoder='ROAD adapted' if d['condition']!='frozen' else 'original',seed=0,split='420 expert /90 development',data=d,source=str(p.relative_to(W))))
runrows=[['Study','Model','Encoder','Seed','Metric','Evaluation frames',*heads,*tails,'Training split','Source']]
perclass=[['Study','Model','Encoder','Seed','Head','Class index (0-based)','Class label','AP (%)','Source']]
groups=defaultdict(list)
for r in records:
 d=r['data'];vals=[d['summary'].get(h) for h in heads]+[d.get('tail',{}).get(t,{}).get('mAP') for t in tails]
 runrows.append([r['study'],r['model'],r['encoder'],r['seed'],'Detector frame-mAP @ IoU0.5 (%)',d.get('n_frames'),*vals,r['split'],r['source']]);groups[(r['study'],r['model'],r['encoder'],r['split'])].append((vals,r['source']))
 for h,arr in d.get('ap_values',{}).items():
  lines=d.get('per_class',{}).get(h,[])
  for i,v in enumerate(arr):
   label=lines[i].split(' : ')[0] if i<len(lines) else None
   perclass.append([r['study'],r['model'],r['encoder'],r['seed'],h,i,label,v,r['source']])
agg=[['Study','Model','Encoder','Runs','Metric']+[x for h in heads+tails for x in [h+' mean',h+' sample SD']]+['Training split','Sources']]
for k,rows in groups.items():
 values=[]
 for i in range(9):
  a=[r[0][i] for r in rows if r[0][i] is not None];values.extend([statistics.mean(a) if a else None,statistics.stdev(a) if len(a)>1 else None])
 agg.append([*k[:3],len(rows),'Detector frame-mAP @ IoU0.5 (%)',*values,k[3],'; '.join(r[1] for r in rows)])
bdd=[agg[0]]+[r for r in agg[1:] if r[0]=='BDD-X paired']
crop=[['Study','Stage','Condition','Checkpoint','Metric','Triplet crop AP (%)','Selection so far','Training scope','Source']]
for p in sorted((W/'artifacts/stage56-full/collected').glob('*.partial.json')):
 d=read(p)
 for label,m in [('baseline',d['baseline'])]+[(f"epoch-{e['epoch']}",e) for e in d['epochs']]:
  crop.append(['ROAD full adaptation',d['stage'],d['condition'],label,'Development triplet crop AP; NOT detector AP',m['triplet_crop_AP'],d['selected'],'420 train /90 development; seed0',str(p.relative_to(W))])
p=W/'artifacts/road-contrastive-adaptation/pilot-summary.json'
for d in read(p):
 for i,ap in enumerate([d['baseline']]+d['epoch_AP']):crop.append(['ROAD small pilot','Flat Linear184 ONLY',d['condition'],'baseline' if i==0 else f'epoch-{i}','Subset development triplet crop AP; NOT detector AP',ap,d['selected'],'941 training rows /247 development rows; 31 of86 dev triplets',str(p.relative_to(W))])
p=W/'artifacts/bddx-contrastive-pilot/collected/run-seed0/report.json';d=read(p)
retrieval=[['Dataset','Checkpoint','Development contrastive loss','Video-to-text R@1 (%)','Text-to-video R@1 (%)','Train videos','Dev videos','Selected','Source']]
for label,m in [('baseline',d['baseline'])]+[(f"epoch-{e['epoch']}",e) for e in d['epochs_results']]:retrieval.append(['BDD-X subset',label,m['loss'],100*m['v2t_r1'],100*m['t2v_r1'],d['train_n'],d['dev_n'],d['selected'],str(p.relative_to(W))])
status=[['Study','Stage','Condition','Epoch','Frames done','Frames total','Optimizer updates','Last progress UTC','Status / limitations']]
for p in sorted((W/'artifacts/stage56-full/collected').glob('*.progress.json')):
 d=read(p);status.append(['ROAD full adaptation',d['stage'],d['condition'],d['epoch'],d['frames_done'],d['frames_total'],d['steps'],d['updated_utc'],'Progress snapshot only; use completed crop reports for scores'])
p=W/'artifacts/stage56-full/collected/epoch1-launch.json'
if p.exists():
 d=read(p)
 for k in ['stage5_adaptation','controls','stage6_adaptation']:status.append(['Epoch1 detector evaluation',k,None,None,None,None,None,d['submitted_utc'],'Submitted job '+d[k]+'; do not interpret absence of result as zero AP'])
frcb=[['Reference','Dataset','Metric','Runs','Agent','Action','Location','Duplex','Triplet','Triplet SD','Notes','Source'],['FRCB Table5','ROAD-Waymo','frame-mAP @IoU0.5 (%)',5,33.49,23.28,31.66,17.58,11.53,.35,'External published protocol; not matched candidates/frames; tail membership not aligned','https://doi.org/10.1016/j.patcog.2026.114489'],['FRCB Table3','ROAD-Waymo','frame-mAP @IoU0.5 (%)','best-result',33.64,23.87,32.33,17.84,11.71,None,'Keep separate from mean across runs','wiki/papers/zhong-2026-frcb.md']]
notes=[['Topic','Record'],['Updated UTC',now],['Scope','New results since original workbook: controlled original encoders, paired BDD-X, BCE gating, AP-selected blending, small ROAD pilot, full ROAD epoch reports and FRCB reference.'],['Preservation','Original worksheets, drawings, images, relationships and cell values retained unchanged. New sheets appended.'],['Metric separation','Detector frame-mAP, development crop AP and BDD-X retrieval have separate tabs. Never directly rank crop AP against detector AP.'],['Tail groups','tail47: z<0; deep28: z<-0.5 (subset of tail47); common39: z>=0. Training-frequency groups from each locked study.'],['Aggregation','Means and sample SD computed from raw per-seed detector reports. Blank SD means single result, not zero variance.'],['Comparisons','Compare within study and matched protocol. BDD-X paired original differs from original controlled study due to paired recaching. Blend studies use420 expert videos.'],['BDD-X scope','1024 train /128 dev videos, one adaptation seed,3 epochs; epoch2 selected by BDD-X dev contrastive loss. Three downstream seeds measure head variability only.'],['BDD-X conclusion','No consistent downstream benefit: Stage5 roughly unchanged, Stage6 slightly lower; full-data BDD-X untested.'],['ROAD full scope','Stage5/6 x frozen/classification/contrastive. Last visual block adapts; text frozen. Focal gamma2 over184 channels plus0.001 triplet contrastive in contrastive arms.'],['Classes','184 output channels include agentness +183 semantic labels. Triplet crop AP macro average over86 classes;84 present in full development and absent classes counted zero.'],['Early results','Stage5 epoch1 contrastive > classification > frozen, one seed; no significance or tail gain established from overall crop AP alone.'],['Pending','No placeholder numeric zeros for unfinished Stage6 epochs or detector evaluations. Status tab records latest available progress timestamps.'],['Deadline','Matched study Sep25; committee draft Oct1; defense Oct14-16.'],['New results destination','This local master workbook is the user-designated destination. No SharePoint workbook updated.'],['Sources','Each numeric row references raw report path relative to /data/repos/wiki; source hashes in wiki/artifacts/results-master/.']]
sheets={'New results guide':notes,'Detector summary':agg,'Detector runs':runrows,'BDD-X paired summary':bdd,'Development crop AP':crop,'BDD-X retrieval':retrieval,'Training status':status,'Published FRCB':frcb,'New per-class detector AP':perclass}
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships';PKG='http://schemas.openxmlformats.org/package/2006/relationships';CT='http://schemas.openxmlformats.org/package/2006/content-types'
orig=target.read_bytes();backup=out/f'road_waymo_results_master-before-{stamp}.xlsx';backup.write_bytes(orig)
with zipfile.ZipFile(target) as z:entries={n:z.read(n) for n in z.namelist()}
wb=E.fromstring(entries['xl/workbook.xml']);rels=E.fromstring(entries['xl/_rels/workbook.xml.rels']);cts=E.fromstring(entries['[Content_Types].xml']);slist=wb.find('{'+NS+'}sheets');existing={s.get('name') for s in slist};assert not existing.intersection(sheets),'New tabs already exist; explicit update required'
maxid=max(int(s.get('sheetId')) for s in slist);relids={r.get('Id') for r in rels};newentries={}
stylewb=openpyxl.load_workbook(target);headerstyle=stylewb.worksheets[0]['A1'].style_id
for offset,(name,rows) in enumerate(sheets.items(),1):
 sid=maxid+offset;rid=f'rIdNewResults{sid}';assert rid not in relids
 part=f'xl/worksheets/newresults{sid}.xml'
 node=E.Element('{'+NS+'}worksheet',nsmap={None:NS});views=E.SubElement(node,'{'+NS+'}sheetViews');view=E.SubElement(views,'{'+NS+'}sheetView',workbookViewId='0');E.SubElement(view,'{'+NS+'}pane',ySplit='1',topLeftCell='A2',activePane='bottomLeft',state='frozen');cols=E.SubElement(node,'{'+NS+'}cols')
 for c in range(1,len(rows[0])+1):E.SubElement(cols,'{'+NS+'}col',min=str(c),max=str(c),width=str(30 if c<=3 else 19),customWidth='1')
 if name=='New results guide':cols[-1].set('width','115')
 data=E.SubElement(node,'{'+NS+'}sheetData')
 for ri,row in enumerate(rows,1):
  rn=E.SubElement(data,'{'+NS+'}row',r=str(ri))
  for ci,v in enumerate(row,1):
   if v is None:continue
   cn=E.SubElement(rn,'{'+NS+'}c',r=f'{openpyxl.utils.get_column_letter(ci)}{ri}')
   if ri==1:cn.set('s',str(headerstyle))
   if isinstance(v,(int,float)):E.SubElement(cn,'{'+NS+'}v').text=str(v)
   else:cn.set('t','inlineStr');E.SubElement(E.SubElement(cn,'{'+NS+'}is'),'{'+NS+'}t').text=str(v)
 E.SubElement(node,'{'+NS+'}autoFilter',ref=f'A1:{openpyxl.utils.get_column_letter(len(rows[0]))}{len(rows)}')
 newentries[part]=E.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
 E.SubElement(slist,'{'+NS+'}sheet',name=name,sheetId=str(sid),attrib={'{'+REL+'}id':rid})
 E.SubElement(rels,'{'+PKG+'}Relationship',Id=rid,Type=REL+'/worksheet',Target='worksheets/'+part.rsplit('/',1)[1])
 E.SubElement(cts,'{'+CT+'}Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
newentries.update({'xl/workbook.xml':E.tostring(wb,xml_declaration=True,encoding='UTF-8'), 'xl/_rels/workbook.xml.rels':E.tostring(rels,xml_declaration=True,encoding='UTF-8'),'[Content_Types].xml':E.tostring(cts,xml_declaration=True,encoding='UTF-8')})
tmp=target.with_name(target.stem+'.updating.xlsx')
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in {**entries,**newentries}.items():z.writestr(n,b)
check=openpyxl.load_workbook(tmp)
for before in stylewb:
 after=check[before.title];assert list(before.values)==list(after.values);assert len(before._images)==len(after._images)
for name,rows in sheets.items():
 assert list(check[name].values)==[tuple(row+[None]*(len(rows[0])-len(row))) for row in rows],name
with zipfile.ZipFile(tmp) as z:
 for n,b in entries.items():
  if n not in newentries:assert z.read(n)==b,n
assert target.read_bytes()==orig,'Workbook changed during update'
os.replace(tmp,target)
manifest={'updated_utc':now,'workbook':str(target),'backup':str(backup),'before_sha256':hashlib.sha256(orig).hexdigest(),'after_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'sheets':{k:len(v)-1 for k,v in sheets.items()},'source_sha256':sources,'validation':'All pre-existing cells/images and unmodified ZIP parts preserved; every new cell read back exactly'}
(out/f'update-{stamp}.json').write_text(json.dumps(manifest,indent=2)+'\n');print(json.dumps({k:v for k,v in manifest.items() if k!='source_sha256'},indent=2))
