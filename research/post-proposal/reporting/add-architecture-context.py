from pathlib import Path
from copy import deepcopy
from lxml import etree as E
import openpyxl,zipfile,datetime,json,hashlib,os
p=Path('/home/brandon/Downloads/ROAD-Waymo Results.xlsx');w=openpyxl.load_workbook(p);s=w.worksheets[0];N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda t:'{'+N+'}'+t
specific={3:'Published spatiotemporal detection reference.',4:'Published alternative video backbone reference.',5:'Reference for actor localization and agent classification only.',6:'Published tube-detection reference; different metric and split.',8:'Establish the locally reproduced detection baseline.',9:'Test logical consistency penalties; this historical constraint result is tainted.',10:'Test whether VLM features support dense detection.',11:'Test query-based detection from VLM features.',12:'Test a conventional CNN backbone with deformable attention.',13:'Test whether frozen CLIP context improves detection.',14:'Test whether increasing backbone capacity resolves weak detection.',15:'Test whether input resolution explains localization weakness.',16:'Test flat multilabel outputs with explicit negative-query supervision.',17:'Test whether extra positive query assignments improve detector training.',19:'Hold boxes fixed to isolate scoring improvements.',20:'Test zero-shot language-model classification of detected actors.',21:'Test whether detector priors improve VLM predictions.',22:'Test whether task/language tuning improves structured text predictions.',23:'Test whether language evidence adds to detector scores.',24:'Check whether fusion gains persist when language is removed.',25:'Test whether BDD-X tuning makes VLM evidence more useful.',26:'Test whether joint task/language tuning improves fusion.',27:'Test task heads on visual representations instead of text-output predictions.',28:'Test corrected logical constraints at low weight.',29:'Test corrected logical constraints at medium weight.',30:'Test corrected logical constraints at high weight.',31:'Test whether joint language supervision helps ROAD perception.',32:'Test combined language supervision and logical constraints.',33:'Isolate the change to a vision-language-pretrained encoder.',34:'Test class phrases as classifier weights.',37:'Improve actor proposals before downstream action/composition scoring.',38:'Test a newer detector under the ROAD recipe; proposal-depth caveat applies.',41:'Test whether logic loss helps the RoI classifier.',42:'Test whether detector-derived negatives improve false-positive ranking.',43:'Test parameter-free composition against learned composition outputs.',44:'Test whether primitive scores alone support learned compositions.',46:'Test whether more optimizer updates explain score gains.',48:'Isolate encoder replacement before adding phrase classification.',49:'Test phrase-based classification on the same full-frame features.',50:'Test whether phrase classification requires pretrained language alignment.',51:'Test learned compositions from phrase primitives and I3D features.',53:'Check whether probe-trained crop heads generalize to full validation.',54:'Check whether the probe-trained phrase head generalizes to full validation.',55:'Test full-data learning from actor-focused crop features.',58:'Test phrase primitives plus crop features as composition inputs.',59:'Same historical configuration as preceding row; retained duplicate, not independent evidence.',60:'Check sensitivity to the class-phrase wording.',62:'Provide the matched control for adding box coordinates.',63:'Test whether explicit box position/size helps classification.',65:'Separate extra score-input capacity from phrase semantics.',66:'Test whether correct phrase-to-class correspondence matters.',69:'Test whether gains depend on correct class-phrase alignment.',77:'Test whether BDD-X adaptation improves downstream visual composition.',78:'Test whether BDD-X adaptation improves phrase-fusion transfer.',79:'Matched original-encoder control for BDD-X adaptation.',80:'Matched original-encoder control for BDD-X phrase-fusion transfer.',82:'Learn class-specific reliance on phrase versus visual scores using BCE.',83:'Learn one global phrase/visual mixture weight using BCE.',84:'Flat-head reference using the same reduced expert-training split.',85:'Phrase-head reference using the same reduced expert-training split.',86:'Check semantic dependence of the class-specific gate.',87:'Check semantic dependence of the BCE global blend.',88:'Visual-composition reference on the matched gate-study split.',89:'Phrase-fusion reference on the matched gate-study split.',91:'Select a global mixture for development AP ranking rather than BCE fit.',92:'Check whether the AP-selected gain depends on correctly assigned phrase evidence.',120:'Published visual/context and class-balancing comparison.',121:'Published best-result reference, distinct from its five-run average.'}
def purpose(ri,name,variant):
 if 102<=ri<=113:
  if 'Pre-adaptation' in variant:return 'Small-pilot starting reference: flat184 head only; no Comp MLP and not Stage5/6.'
  if 'Frozen encoder' in variant:return 'Small pilot: isolate head-only training; flat184 head, no Comp MLP, not Stage5/6.'
  if 'contrastive' in variant:return 'Small pilot: test added phrase-alignment supervision before full Stage5/6 runs; no Comp MLP.'
  return 'Small pilot: test task-specific encoder gradients through a flat184 head before full Stage5/6 runs; no Comp MLP.'
 if 94<=ri<=101:
  if 'Pre-adaptation' in variant:return 'Starting checkpoint for this arm; measure change from continued training.'
  if 'Frozen encoder' in variant:return 'Isolate gains from further head training without changing crop representations.'
  if 'contrastive' in variant:return 'Test whether triplet phrase alignment adds beyond focal-only encoder adaptation.'
  return 'Test whether ROAD focal supervision improves task-specific crop representations.'
 if 115<=ri<=118:return 'Reference for video-text retrieval before adaptation.' if ri==115 else 'Test driving-domain video-text alignment on a BDD-X subset before ROAD transfer.'
 if ri in specific:return specific[ri]
 for prefix,text in [('Stage 0','Provide a matched detector baseline.'),('Stage 1:','Reuse stronger YOLO proposals while transferring existing I3D scores.'),('Stage 2:','Score each YOLO box directly from RoI features instead of relying on detection matching.'),('Stage 3:','Learn duplex/triplet interactions from primitive scores and I3D features.'),('Stage 4:','Use actor-focused crop features and class phrases to test semantic classification.'),('Stage 5:','Combine flat primitive evidence and crop features in a learned composition MLP.'),('Stage 6:','Test whether phrase-composition scores add beyond Stage5 visual composition.'),('Crop Flat Head','Provide a crop-feature classifier without explicit phrase scores or composition MLP.')]:
  if name.startswith(prefix):return text
 raise ValueError((ri,name))
raw=p.read_bytes();stamp=datetime.datetime.now().strftime('%Y%m%d-%H%M%S');out=Path('/data/repos/wiki/artifacts/results-master');backup=out/f'ROAD-Waymo-Results-before-context-{stamp}.xlsx';backup.write_bytes(raw)
with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
r=E.fromstring(parts['xl/worksheets/sheet1.xml']);st=E.fromstring(parts['xl/styles.xml']);xfs=st.find(q('cellXfs'));stylemap={};changes=[]
for row in r.find(q('sheetData')):
 ri=int(row.get('r'))
 if ri==1 or s.cell(ri,1).value is None:continue
 name=s.cell(ri,2).value;old=s.cell(ri,3).value or '';why=purpose(ri,name,old);value=old+'\nPurpose: '+why;c=next(c for c in row if c.get('r')==f'C{ri}')
 for x in list(c):c.remove(x)
 c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=value
 original=int(c.get('s','0'))
 if original not in stylemap:
  xf=deepcopy(xfs[original]);a=xf.find(q('alignment'))
  if a is None:a=E.SubElement(xf,q('alignment'))
  a.set('wrapText','1');a.set('vertical','top');xf.set('applyAlignment','1');stylemap[original]=len(xfs);xfs.append(xf)
 c.set('s',str(stylemap[original]));row.set('ht',str(max(float(row.get('ht','15')),110)));row.set('customHeight','1');changes.append({'row':ri,'variant':old,'purpose':why})
xfs.set('count',str(len(xfs)))
for col in r.find(q('cols')):
 if col.get('min')=='3':col.set('width','48')
updates={'xl/worksheets/sheet1.xml':E.tostring(r),'xl/styles.xml':E.tostring(st)};tmp=p.with_name('ROAD-Waymo Results.updating.xlsx')
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in {**parts,**updates}.items():z.writestr(n,b)
a=openpyxl.load_workbook(tmp)
for row in s:
 for c in row:
  if c.column!=3:assert a.worksheets[0].cell(c.row,c.column).value==c.value,c.coordinate
for change in changes:
 c=a.worksheets[0].cell(change['row'],3);assert change['purpose'] in c.value and c.alignment.wrap_text
with zipfile.ZipFile(tmp) as z:
 for n,b in parts.items():
  if n not in updates:assert z.read(n)==b
assert p.read_bytes()==raw;os.replace(tmp,p);(out/f'context-{stamp}.json').write_text(json.dumps({'file':str(p),'backup':str(backup),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'rows':changes,'validation':'Only Variant text, Variant wrapping/width and row heights changed; other cell values and diagram bytes exact.'},indent=2));print('Added purpose statements to',len(changes),'variants; pilot scope explicit; all numerical cells unchanged.')
