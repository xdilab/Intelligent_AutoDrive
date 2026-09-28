"""Read-only aggregation of contextual evolution evidence; no single-seed selection."""
from pathlib import Path
import hashlib, json, statistics
from openpyxl import load_workbook

ROOT=Path('/data/repos/wiki')
OUT=ROOT/'artifacts/contextual-evolution-20260924'
BOOK=Path('/home/brandon/Downloads/ROAD-Waymo Results(1).xlsx')
families={
 88:('contextual-roi-20260917','mlp-classification'),
 89:('contextual-roi-20260917','attention-classification'),
 90:('contextual-roi-20260917','mlp-contrastive'),
 91:('contextual-roi-20260917','attention-contrastive'),
 92:('contextual-roi-all184-20260918','mlp-contrastive-all184'),
 93:('contextual-roi-all184-20260918','attention-contrastive-all184'),
 94:('contextual-gate-20260919','global'),
 95:('contextual-gate-20260919','class'),
 96:('contextual-gate-20260919','generic'),
 97:('contextual-gate-20260919','factorized'),
 102:('contextual-roi-dcb-20260921','attention-classification-dcb'),
 103:('contextual-roi-dcb-20260921','attention-contrastive-all184-dcb'),
 104:('contextual-roi-dcb-20260921','attention-global-dcb'),
}
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
sheet=load_workbook(BOOK,read_only=True,data_only=True).worksheets[0]
rows={r[0]:r for r in sheet.iter_rows(values_only=True) if isinstance(r[0],int)}
metrics=['agentness','agent','action','loc','duplex','triplet','tail47','deep28','common39']
result={}
for ident,(study,condition) in families.items():
 p=ROOT/'artifacts'/study/'all-results.json'
 source=json.loads(p.read_text())
 runs=[r for r in source['runs'] if r.get('kind')==condition or r.get('run','').rsplit('-seed',1)[0]==condition]
 assert len(runs)==3
 vals={k:[r['summary'][k] if k in r['summary'] else r['tail'][k]['mAP'] for r in runs] for k in metrics}
 means={k:statistics.mean(v) for k,v in vals.items()}
 for k,col in zip(metrics[:7],range(12,19)):
  assert abs(means[k]-rows[ident][col])<1e-10,(ident,k,means[k],rows[ident][col])
 result[ident]={'id':ident,'name':rows[ident][1],'condition':condition,'source':str(p),'source_sha256':sha(p),'n':3,'mean':means,'sample_sd':{k:statistics.stdev(v) for k,v in vals.items()},'runs':[{'run':r.get('run',f"{r.get('kind')}-seed{r.get('seed')}"),'summary':r['summary'],'tail':r['tail'],'selected_epoch':r.get('selected_epoch'),'signature':r['signature']} for r in runs]}
provenance=[BOOK]
for study,_ in families.values():
 provenance.extend([ROOT/'artifacts'/study/'all-results.json',ROOT/'artifacts'/study/'protocol.json'])
repo=Path('/data/repos/ROAD_Reason/research/post-proposal')
for p in ['contextual-roi/model.py','contextual-roi/cache_context.py','contextual-roi/train_cached.py','contextual-roi/evaluate_cached.py','contextual-roi-all184/model.py','contextual-roi-dcb/model.py','contextual-roi-dcb/dcb.py','contextual-roi-dcb/prepare_blend.py','contextual-roi-dcb/protocol.json','contextual-gate/router.py','contextual-gate/common.py','contextual-gate/train_gate.py','stage56-full/metric.py']:
 provenance.append(repo/p)
out={'selected_diagram_ids':[88,89,90,92,94,104],'selection_note':'Condition means over all three seeds. #94 is the simple global-blend lineage ancestor of #104; #97 narrowly wins triplet among gates and must be disclosed. MLP contrastive representatives #90/#92 are not parents of attention blend #94.','workbook_sha256':sha(BOOK),'metrics':metrics,'models':result,'sources':{str(p):sha(p) for p in sorted(set(provenance))}}
(OUT/'selection.json').write_text(json.dumps(out,indent=2)+'\n')
lines=['| ID | Condition | '+' | '.join(metrics)+' |','|---|---|'+'---:|'*len(metrics)]
for ident,d in result.items():lines.append('| '+str(ident)+' | '+d['condition']+' | '+' | '.join(f'{d["mean"][k]:.5f}' for k in metrics)+' |')
(OUT/'metrics-table.md').write_text('\n'.join(lines)+'\n')
print('\n'.join(lines))
print('Verified workbook matches all 13 condition means; wrote selection.json with source hashes.')
