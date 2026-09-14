"""Import every nonempty workbook row without treating research notes as verified facts."""
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import csv, hashlib, html, json, re, runpy

ROOT=Path(__file__).resolve().parent
SOURCE=Path('/home/brandon/Downloads/_Papers Abstracts_BrandonByrd.xlsx')
def norm(s):return re.sub(r'[^a-z0-9]','',str(s).lower())
def urls(s):return re.findall(r'https?://[^\s<>"\]]+',str(s or ''))
specs={
 'Papers':(3,2,1,None,None),
 'Pretrained Model Papers':(5,None,6,None,12),
 'Pedestrian Intention Lit Review':(2,None,6,None,13),
 'ROAD Lit Review':(1,4,3,2,None),
 'Papers to look at ':(1,None,2,None,8),
 'Pretrained Resources':(2,None,None,None,4),
}
w=load_workbook(SOURCE,data_only=False);raw=[];records=[];notes=[];allurls=[]
for s in w:
 for row in s.iter_rows(min_row=2):
  cells=[{'column':c.column,'value':str(c.value),'hyperlink':c.hyperlink.target if c.hyperlink else None} for c in row if c.value is not None or c.hyperlink]
  if not cells:continue
  loc=f'{s.title}!{row[0].row}'
  raw.append({'sheet':s.title,'row':row[0].row,'cells':cells})
  for c in cells:
   for u in urls(c['value'])+([c['hyperlink']] if c['hyperlink'] else []):allurls.append({'source':loc,'column':c['column'],'url':u})
  if s.title not in specs:notes.append({'source':loc,'reason':'Metric/source note','cells':cells});continue
  tc,yc,vc,ac,pc=specs[s.title]
  def v(c):return str(row[c-1].value or '').strip() if c else ''
  title=v(tc)
  if not title or title in ['Road cited','Road -R cited','Road-waymo cited'] or title.startswith('How they help'):
   notes.append({'source':loc,'reason':'Search pointer or research question, not a paper','cells':cells});continue
  year=v(yc) or (re.search(r'20\d\d',v(vc)).group() if re.search(r'20\d\d',v(vc)) else '')
  paperurls=urls(v(pc)) if pc else []
  rowurls=[x['url'] for x in allurls if x['source']==loc]
  # A generic model-zoo link is not sufficient to merge two different papers.
  paperurls=paperurls or [u for u in rowurls if any(x in u for x in ['arxiv.org/','openaccess.thecvf.com/','sciencedirect.com/science/','openreview.net/'])]
  records.append({'source':loc,'title':title,'year':year,'venue':v(vc),'author':v(ac).replace(';',' and '),'url':paperurls[0].rstrip('.,') if paperurls else (rowurls[0].rstrip('.,') if rowurls else ''),'all_urls':rowurls})
ROOT.joinpath('workbook-source.json').write_text(json.dumps({'path':str(SOURCE),'sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'sheets':w.sheetnames,'rows':raw},indent=2,ensure_ascii=False)+'\n')
ROOT.joinpath('workbook-records.json').write_text(json.dumps(records,indent=2,ensure_ascii=False)+'\n')
ROOT.joinpath('workbook-notes.json').write_text(json.dumps(notes,indent=2,ensure_ascii=False)+'\n')
ROOT.joinpath('workbook-link-ledger.json').write_text(json.dumps(allurls,indent=2,ensure_ascii=False)+'\n')

if __name__=='__main__':
 print(json.dumps({'sheets':len(w.sheetnames),'nonempty_rows':len(raw),'citation_rows':len(records),'note_rows':len(notes),'url_mentions':len(allurls)}))

# Rebuild the previously verified core, then extend it from the workbook.
core=runpy.run_path(str(ROOT/'build-bibliography.py'))
base=json.loads((ROOT/'citation-inventory.json').read_text())
meta=json.loads((ROOT/'workbook-primary-metadata.json').read_text())
matches=json.loads((ROOT/'workbook-arxiv-matches.json').read_text())
manual={}
def checked(source,title,author,year,venue,url,kind='inproceedings',**extra):
 manual[source]=dict(title=title,author=author,year=str(year),venue=venue,url=url,kind=kind,**extra)
checked('ROAD Lit Review!2','CCN+: A neuro-symbolic framework for deep learning with requirements','Giunchiglia, Eleonora and Tatomir, Alex and Stoian, Mihaela Cătălina and Lukasiewicz, Thomas',2024,'International Journal of Approximate Reasoning','https://doi.org/10.1016/j.ijar.2024.109124','article',doi='10.1016/j.ijar.2024.109124',volume='171',pages='109124')
checked('Pretrained Model Papers!5','YOLOPX: Anchor-free multi-task learning network for panoptic driving perception','Zhan, Jiao and Luo, Yarong and Guo, Chi and Wu, Yejun and Meng, Jiawei and Liu, Jingnan',2024,'Pattern Recognition','https://doi.org/10.1016/j.patcog.2023.110152','article',doi='10.1016/j.patcog.2023.110152',volume='148',pages='110152')
checked('Pretrained Model Papers!7','ISP-Teacher: Image Signal Process with Disentanglement Regularization for Unsupervised Domain Adaptive Dark Object Detection','Zhang, Yin and Zhang, Yongqiang and Zhang, Zian and Zhang, Man and Tian, Rui and Ding, Mingli',2024,'Proceedings of the AAAI Conference on Artificial Intelligence','https://ojs.aaai.org/index.php/AAAI/article/view/28569',doi='10.1609/aaai.v38i7.28569',volume='38',number='7',pages='7387--7395')
checked('Pretrained Model Papers!17','Technical Report for CVPR 2022 Workshop on Autonomous Driving: Argoverse 3D Object Detection Competition','Fang, Jin and Meng, Qinghao and Zhou, Dingfu and Tang, Chulin and Shen, Jianbing and Xu, Cheng-Zhong and Zhang, Liangjun',2022,'Workshop technical report','https://www.argoverse.org/assets/pdfs/Detectors_Argoverse_2022.pdf','misc')
checked('Pretrained Model Papers!26','3D-AWARE (KITTI benchmark submission)','{KITTI Benchmark}',2025,'Benchmark result page','https://www.cvlibs.net/datasets/kitti/eval_object_detail.php?result=bdcb4d589626c4faf5a63d2cd272ac2d7fa781c0','misc',note='Official indexed entry identifies 3D-AWARE and submitter Fazal Ghaffar, July 19, 2025. Cited as a benchmark webpage; no standalone paper venue or full paper author list is asserted.')
checked('Papers to look at !4','AI vs. Humans: Comparing road user intention recognition performance','Vellenga, Koen and Steinhauer, H. Joe and Falkman, Göran and Andersson, Jonas and Sjögren, Anders',2026,'Transportation Research Part F: Traffic Psychology and Behaviour','https://doi.org/10.1016/j.trf.2025.103491','article',doi='10.1016/j.trf.2025.103491',volume='118',pages='103491',note='2025 DOI/online record; final volume dated March 2026.')
checked('Papers to look at !6','T-norm Selection for Object Detection in Autonomous Driving with Logical Constraints','Eiter, Thomas and Higuera Ruiz, Nelson and Inoue, Katsumi and Moriyama, Sota',2025,'Advances in Neural Information Processing Systems','https://papers.nips.cc/paper_files/paper/2025/file/7dbdf006424e7749c8a35913d3574c4e-Paper-Conference.pdf')
checked('ROAD Lit Review!37','ROAD-R 2023: the Road Event Detection with Requirements Challenge','Giunchiglia, Eleonora and Stoian, Mihaela C. and Khan, Salman and Alitappeh, Reza Javanmard and Teeti, Izzeddin A. M. and Paschke, Adrian and Cuzzolin, Fabio and Lukasiewicz, Thomas',2023,'NeurIPS competition website','https://neurips.cc/virtual/2023/competition/66596','misc')
checked('Papers to look at !2:linked-chapter','Towards a VLM-Based Foundation for Generalised Neurosymbolic Visual Commonsense','Suchan, Jakob and Baloch, Salim and Bhatt, Mehul',2026,'Foundations of Information and Knowledge Systems (FoIKS), Lecture Notes in Computer Science','https://doi.org/10.1007/978-3-032-21540-6_24',volume='16475',pages='359--365',doi='10.1007/978-3-032-21540-6_24')
for source,title,author,url in [
 ('Pretrained Model Papers!3','BDD100K Model Zoo','{SysCV}','https://github.com/SysCV/bdd100k-models'),
 ('Pedestrian Intention Lit Review!3','Joint Attention in Autonomous Driving (JAAD): Dataset and annotations','{JAAD Dataset Authors}','https://github.com/ykotseruba/JAAD'),
 ('Pretrained Resources!2','OpenPCDet Model Zoo','{OpenPCDet Contributors}','https://github.com/open-mmlab/OpenPCDet'),
 ('Pretrained Resources!4','RayDN official repository','{RayDN Authors}','https://github.com/LiewFeng/RayDN'),
 ('Pretrained Resources!5','BEVFusion official repository','{BEVFusion Authors}','https://github.com/mit-han-lab/bevfusion')]:
 checked(source,title,author,2026,'Dataset/software resource',url,'misc',note='Accessed September 7, 2026; access year is not a release-date claim.')

def canonical_url(u):
 if 'arxiv.org' in u:
  m=re.search(r'(\d{4}\.\d{4,5})',u)
  if m:return 'https://arxiv.org/abs/'+m[1]
 if 'openaccess.thecvf.com' in u:return u.replace('/papers/','/html/').replace('_paper.pdf','_paper.html')
 return u.rstrip('/')

titlemeta={norm(e['title']):e for e in meta.values()}
aliases={norm('ROAD-Waymo: Action Awareness at Scale for Autonomous Driving'):'khan2024roadwaymo'}
bytitle={norm(e['title']):e for e in base};byurl={canonical_url(e['url']):e for e in base}
bykey={e['key']:e for e in base}
for e in base:e['verification']='Primary source checked in core audit';e['workbook_sources']=[]
mapping=[];corrections=[]
records.append(dict(source='Papers to look at !2:linked-chapter',title=manual['Papers to look at !2:linked-chapter']['title'],url=manual['Papers to look at !2:linked-chapter']['url'],author='',year='',venue='',all_urls=[]))
for r in records:
 u=canonical_url('https://arxiv.org/abs/'+matches[r['source']] if r['source'] in matches else r['url'])
 m=manual.get(r['source']) or meta.get(u) or titlemeta.get(norm(r['title']))
 if r['source']=='Pretrained Resources!3':m=manual['Pretrained Model Papers!3']
 e=bykey.get(aliases.get(norm(r['title']))) or bytitle.get(norm(r['title']))
 if not e and m:e=bytitle.get(norm(m['title'])) or byurl.get(canonical_url(m['url']))
 # Exact paper identifiers support merging shorthand duplicate rows.
 if not e and ('arxiv.org' in u or 'openaccess.thecvf.com' in u):e=byurl.get(u)
 if e:
  e['workbook_sources'].append(r['source']);mapping.append(dict(source=r['source'],key=e['key'],disposition='merged',original_title=r['title']));continue
 if m:
  e=dict(m);e['verification']='Primary bibliographic metadata checked'
  if 'arxiv.org' in e['url']:
   e['eprint']=e['url'].split('/')[-1];e['kind']='misc';e['venue']='arXiv preprint arXiv:'+e['eprint']
   e['note']='Cites the verified arXiv version and first-submission year; a later published venue may also exist. Workbook venue/year retained in row provenance, not silently accepted.'
  else:
   e.setdefault('kind','inproceedings')
   if 'CVPR2025W' in e['url']:e['venue']='Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition Workshops'
  if r['source'].startswith('Pretrained Resources') or r['source'] in ['Pretrained Model Papers!3','Pretrained Model Papers!26','Pedestrian Intention Lit Review!3']:e['verification']='Official resource identity checked'
  for k in ['title','author','year','venue']:
   if r.get(k) and norm(r[k])!=norm(e.get(k,'')):corrections.append(dict(source=r['source'],field=k,workbook=r[k],canonical=e.get(k,''),evidence=e['url']))
 else:
  e={k:r.get(k,'') for k in ['title','author','year','venue','url']};e['kind']='misc';e['verification']='Unresolved resource identity — do not cite as a paper';e['note']='Workbook record retained. Linked KITTI page did not expose enough entry-specific metadata to authenticate the named method.'
 e['key']='wb_'+re.sub(r'[^a-z0-9]+','_',r['source'].lower()).strip('_')
 e['group']='workbook-background'
 e['role']='Associated reading-list source; inclusion here does not imply implementation, comparison, or support for current thesis results.'
 e['where']=r['source'];e['workbook_sources']=[r['source']]
 base.append(e);bykey[e['key']]=e;bytitle[norm(r['title'])]=e;bytitle[norm(e['title'])]=e;byurl[canonical_url(e['url'])]=e
 mapping.append(dict(source=r['source'],key=e['key'],disposition='added',original_title=r['title']))

# Save row-complete provenance independently of normalized bibliography fields.
ROOT.joinpath('workbook-row-mapping.json').write_text(json.dumps(mapping,indent=2,ensure_ascii=False)+'\n')
ROOT.joinpath('workbook-metadata-corrections.json').write_text(json.dumps(corrections,indent=2,ensure_ascii=False)+'\n')
ROOT.joinpath('citation-inventory.json').write_text(json.dumps(base,indent=2,ensure_ascii=False)+'\n')
cols=['key','group','verification','title','author','year','venue','url','role','where','workbook_sources']
def flat(e):return {k:'; '.join(e[k]) if isinstance(e.get(k),list) else e.get(k,'') for k in cols}
with ROOT.joinpath('citation-inventory.csv').open('w',newline='') as f:
 wr=csv.DictWriter(f,fieldnames=cols);wr.writeheader();wr.writerows(flat(e) for e in base)
md=['# Master thesis-associated bibliography','','Source coverage: current thesis/deck and all seven worksheets of `_Papers Abstracts_BrandonByrd.xlsx`. All 101 nonempty workbook rows are accounted for: 93 paper/resource rows and 8 metric/research notes. One additional paper linked inside a note is also identified. Original workbook contents and 127 URL mentions are retained separately.','',f'**{len(base)} unique paper/resource records.** Core sources are distinguished from associated background reading. The workbook summaries and performance claims have not been adopted as verified facts. ArXiv entries intentionally cite the checked preprint metadata; inventory numbers are not the thesis’s IEEE numbering.','']
bib=['% Full thesis-associated bibliography, including workbook background.','% ArXiv entries cite verified preprint metadata; use keys to select sources actually discussed.','']
for i,e in enumerate(base,1):
 md += [f'## [{i}] {e["title"]}','',f'**Key:** `{e["key"]}` · **Group:** {e["group"]} · **Status:** {e["verification"]}','',f'{core["author_short"](e["author"])}. *{e["venue"]}*, {e["year"]}. [Source]({e["url"]}).','',e['role'],'', '**Placement/provenance:** '+e['where']]
 if e['workbook_sources']:md+=['**Workbook rows:** '+', '.join(e['workbook_sources'])]
 if e.get('note'):md+=['**Note:** '+e['note']]
 md+=['']
 if e['verification'].startswith('Unresolved'):continue
 fields={'author':e['author'],'title':'{'+e['title']+'}','year':e['year']}
 fields['journal' if e['kind']=='article' else 'booktitle' if e['kind']=='inproceedings' else 'howpublished']=e['venue']
 for k in ['volume','number','pages','doi','eprint','url','note']:
  if e.get(k):fields[k]=e[k]
 if e.get('eprint'):fields['archivePrefix']='arXiv'
 fields['keywords']=e['group']
 bib.append('@'+e['kind']+'{'+e['key']+',\n'+',\n'.join('  '+k+' = {'+(v if k in ['url','doi'] else core['bib_escape'](v))+'}' for k,v in fields.items())+'\n}\n')
ROOT.joinpath('master-bibliography.md').write_text('\n'.join(md)+'\n')
ROOT.joinpath('thesis-references.bib').write_text('\n'.join(bib)+'\n')
out=Workbook();s=out.active;s.title='Citation Inventory';s.append(cols)
for e in base:s.append(list(flat(e).values()))
for name,headers,rows in [
 ('Workbook Row Map',['source','key','disposition','original_title'],mapping),
 ('Metadata Corrections',['source','field','workbook','canonical','evidence'],corrections),
 ('Source URLs',['source','column','url'],allurls),
 ('Notes',['source','reason','cells'],notes)]:
 sh=out.create_sheet(name);sh.append(headers)
 for row in rows:sh.append([json.dumps(row.get(k,''),ensure_ascii=False) if isinstance(row.get(k),list) else row.get(k,'') for k in headers])
sh=out.create_sheet('Read Me');sh.append(['Item','Details']);sh.append(['Coverage',f'{len(base)} unique records; all 101 nonempty rows across 7 worksheets accounted for.']);sh.append(['Resources','3D-AWARE is cited as an official benchmark submission, not as an authenticated standalone paper.']);sh.append(['Preprints','ArXiv citations use verified first-submission years; later publication records may differ.']);sh.append(['Claims','Workbook abstracts/results are research notes, not independently validated findings.']);sh.append(['Scope','Background reading does not imply current implementation or endorsement.']);sh.append(['Original source',str(SOURCE)])
for sh in out:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
 for j in range(1,sh.max_column+1):sh.column_dimensions[get_column_letter(j)].width=45 if j>1 else 35
 for row in sh.iter_rows(min_row=2):
  sh.row_dimensions[row[0].row].height=70
  for c in row:
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,str) and c.value.startswith('https://'):c.hyperlink=c.value;c.style='Hyperlink'
out.save(ROOT/'citation-inventory.xlsx')
unresolved=[e for e in base if e['verification'].startswith('Unresolved')]
assert len(mapping)==94 and len({m['source'] for m in mapping})==94
assert len({e['key'] for e in base})==len(base)
assert len(mapping)+len(notes)-1==len(raw)
print(json.dumps({'unique_records':len(base),'bibtex_entries':len(base)-len(unresolved),'row_mappings':len(mapping),'unresolved':[e['title'] for e in unresolved]}))
